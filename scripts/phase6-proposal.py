"""Generate editable/unsubmitted drafts and internal blank-input costing workbook.
Official Exhibit A bytes are copied unchanged; template zeros are never vendor quotes.
"""
from pathlib import Path
import json,hashlib,shutil,importlib.metadata,reportlab
from xml.sax.saxutils import escape
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
root=Path(__file__).resolve().parent.parent;out=root/'artifacts/phase-6';draft=root/'docs/submission'
fresh=json.loads((out/'official-source-refresh.json').read_text());assert fresh['retrieved']
price=next(a for a in fresh['attachments'] if 'Exhibit A' in a['label']);original=root/price['privatePath'];assert hashlib.sha256(original.read_bytes()).hexdigest()==price['sha256']
official=out/'RFP20261404054_Exhibit-A_ORIGINAL-UNFILLED.xlsx';shutil.copyfile(original,official);assert original.read_bytes()==official.read_bytes()
w=load_workbook(original);assert w.sheetnames==['Price Summary'];assert w.active['B12'].value=='=B8+C8+D8+E8'
internal=Workbook();s=internal.active;s.title='Instructions';instructions=[['INTERNAL COST PLANNING — NOT EXHIBIT A / NOT A COMMERCIAL OFFER'],['All actual prices, vendor facts and renewal caps require owner approval. Blank is unknown; zero must never be inferred.'],['Original official sheet is separately preserved byte-for-byte. Its prefilled zeros are template values, not vendor prices.'],['Transfer only authorized final numbers to the original Exhibit A without changing its format.'],['All-inclusive initial three-year term, optional years 4/5, pricing firmness and caps require legal/commercial review.'],['No paid infrastructure, pricing promise, signature, warranty or contract acceptance has been made.']]
for r in instructions:s.append(r)
s.column_dimensions['A'].width=115
s=internal.create_sheet('Cost Inputs');s.append(['Official category','One-time USD','YR1 USD','YR2 USD','YR3 USD','Optional YR4 USD','Optional YR5 USD','Owner evidence / included scope'])
for row in range(3,8):s.append([w.active.cell(row,1).value,None,None,None,None,None,None,'OWNER INPUT REQUIRED'])
s.append(['TOTAL (all five category inputs required)']+[f'=IF(COUNT({c}2:{c}6)=5,SUM({c}2:{c}6),"OWNER INPUT REQUIRED")' for c in 'BCDEFG'])
s.append(['Three-year TCO','=IF(COUNT(B7:E7)=4,SUM(B7:E7),"OWNER INPUT REQUIRED")'])
s.append(['Five-year TCO','=IF(COUNT(B7:G7)=6,SUM(B7:G7),"OWNER INPUT REQUIRED")'])
s.append(['Optional YR4 cap %',None]);s.append(['Optional YR5 cap %',None])
s=internal.create_sheet('Scope Checklist');s.append(['Cost driver','Owner amount/assumption','Inclusion / dependency'])
for x in ['Hosting / HA / backups / DR','Identity / enterprise SSO / MFA','Epic / ADT / interface services','Payer exchange / criteria licenses','Implementation / migration / validation','Clinical governance / physician services','Training / go-live / support staffing','Security / privacy / BAA / assessment','Monitoring / model evaluation / maintenance','Taxes / travel / subcontractors / insurance']:
 s.append([x,None,'OWNER INPUT REQUIRED; no amount or supplier commitment assumed'])
for s in internal:
 s.freeze_panes='B2'
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='243A4A')
 for row in s:
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 if s.title!='Instructions':
  s.column_dimensions['A'].width=42
  for c in 'BCDEFG':s.column_dimensions[c].width=24
  s.column_dimensions['H'].width=65
path=out/'RFP20261404054_INTERNAL-PRICING_OWNER-INPUT.xlsx';internal.save(path)
check=load_workbook(path);assert all(check['Cost Inputs'].cell(r,c).value is None for r in range(2,7) for c in range(2,8));assert all(check['Cost Inputs'].cell(7,c).data_type=='f' for c in range(2,8))
sections=[('DRAFT / DO NOT SUBMIT — RFP #20261404054',[
 'Hospital Utilization Review Technology and Optimization Partner: engineering demonstration and proposed delivery outline.',
 'Respondent legal entity: [OWNER REQUIRED]. Address: [OWNER REQUIRED]. Authorized signatory: [OWNER REQUIRED]. Company experience, staff qualifications and references: [UNVERIFIED OWNER INPUT].',
 'Deadline retained from Addendum 1: November 3, 2026, at 2:00 p.m. CST, as written. Current addenda must be rechecked before any owner-authorized submission.',
 'This draft contains no signature, commercial offer, legal acceptance or production-readiness representation. Missing owner information makes it unsuitable for submission.']),
 ('Executive summary — engineering evidence, not clinical outcomes',[
 'The existing desktop hospital workstation demonstrates all eight requested synthetic workflows: FHIR intake, AI-supported prioritization, inpatient/observation advisory review, continued-stay review, simulated authorization/notification, documentation-risk flags, physician escalation and human review, and reporting/audit.',
 'The modular demonstration reuses Medplum, HAPI FHIR/CQF, Keycloak, PostgreSQL, Grafana, local llama.cpp/Qwen, React and existing OpenTelemetry/Prometheus adapters. No proprietary Epic assets or private criteria were copied. The interface remains desktop-first at 1440/1920 pixels.',
 'Phase 5 retained three actual one-hour queue runs. The final isolated-read-pool result was 265.85 requests/sec and p95 134.14ms, with 45 timeouts and 11 HTTP 502 errors. These are single-host synthetic queue measurements, not clinical inference capacity, an SLA or proof of timeout elimination.',
 'No improvement in clinical accuracy, denials, reimbursement, length of stay or financial ROI is established. Production delivery requires authorized integration, validated clinical/payer criteria, qualified reviewers and security/legal approval.']),
 ('Company background, staffing and references — owner completion',[
 'Supply verified legal entity, history, financial capacity, named implementation/support team, clinical qualifications and comparable hospital delivery record. No team size, certification, customer or experience claim is supplied by the software.',
 'At least three genuine references with authorized name/telephone/email details are required by source line 255. Do not use synthetic patients or test users as customer references.',
 'HUB response must address the 25% good-faith target and applicable current certificate or outreach/selected subcontractor evidence, including historical efforts. No status, outreach, agency contact or subcontract value is invented.']),
 ('Verified solution and limitations',[
 'FHIR R4 fictional intake and version provenance feed persisted cases. Initial and continued-stay review use illustrative CQL plus a pinned local AI model; explanations cite recorded evidence. Missing data remains uncertainty, and recommendations require human review.',
 'Policy Management Center supports evidence/version comparison, correction requests, testing, independent supervisor/clinical-role approval, publication, suspension and rollback. Unapproved changes remain isolated. Rules are illustrative, not licensed medical-necessity or private payer-contract criteria.',
 'Authorization Review displays applicable illustrative documentation requirements, missing evidence and internal human next-step actions. Idempotent local payer simulation preserves receipts and audit. Internal approval is not insurer approval; no real payer submission or patient-status writeback is enabled.',
 'Clinical/payer validation is synthetic engineering evidence only. Real denial prediction, calibrated risk probability, independent sensitivity/specificity, full longitudinal outcomes and complete payer plan coverage remain future work.']),
 ('Proposed integration and implementation — subject to approval',[
 'Proposed gated delivery: owner/hospital discovery and workflow design; approved data/security interfaces; authorized Epic sandbox testing; criteria and payer approvals; independent blinded validation; staged hospital acceptance; role-based training and monitored rollout. No delivery dates or acceptance commitments are made.',
 'Read-only FHIR transport and local SMART/permission/stale-resource/reconciliation tests are engineering demonstrations. Actual Epic tenant/client/scopes and JPS interface permissions are absent. ADT/HL7v2 and certified bidirectional enterprise integration are not established.',
 'Preserve source encounter status, clinical provenance, human decision authority, and explicit approval before external submissions or clinical policy activation. Production requires representative workloads and continuous outcome monitoring.']),
 ('Security, AI governance and operational reliability',[
 'The local synthetic stack has role checks, session boundaries, source lineage, audit hashing, model pinning, monitored fallback, reconciliation and application restore proof. These controls do not establish HIPAA/HITECH compliance, SOC 2/HITRUST/ISO certification or a BAA.',
 'Production gaps include reviewed TLS/encryption at rest and key management, enterprise SSO/MFA, least privilege, privacy risk assessment, penetration testing, retention/incident response, coordinated backup/HA/DR, support staffing and hospital-approved SLAs.',
 'No PHI or JPS production data has been processed. No JPS data is used for external model training. Third-party processing, production model use, fairness/drift monitoring and representative clinical evaluation require separate review.',
 'The Vercel build is a disconnected read-only fictional fixture frontend. The backend stays separate on loopback. No hosted preview or deployment protection is claimed without actual verification.']),
 ('Commercial offer, forms and legal review — not supplied',[
 'Official Exhibit A is preserved unchanged. Its prefilled zeros are not a vendor quote. The internal workbook contains blank actual inputs and conditional TCO formulas; owner must approve all-inclusive prices, initial-term firmness, optional renewals/caps, services, taxes and support assumptions.',
 'Exhibits B (signature), D (vendor certification) and F (good faith) are embedded in the original RFP. C is the separate terms document. E is not utilized. G requires the AI and security/data-governance questionnaires. Owner/counsel must complete/sign as applicable; no attestation or redline acceptance has been made.',
 'All addenda acknowledgments, vendor registration, legal checks, insurance/BAA/criteria agreements, official questionnaire evidence and final review remain outstanding. Form 1295 is post-award as described by the RFP, not an engineering filing task.']),
 ('Response preparation and technical demonstration attachments',[
 'Use the source-exact Phase 6 traceability matrix and vendor-form checklist; engineering statuses are explicitly partial, synthetic, future or owner-dependent. The matrix inventories every retained nonempty source paragraph; it is not a legal opinion or a finding that minimum requirements are satisfied.',
 'Owner must finalize legal entity on every page, cover address/deadline/RFP number, executive summary at most two pages, response at most 25 pages excluding exhibits, authentic three-reference and HUB sections, official Excel pricing and editable/unlocked contract or redline.',
 'The RFP prescribes file upload at https://f1.jpshealth.org/form/RFPResponseForm, RFP number in the Bidding Opportunity Num field, and filenames beginning with RFP number and actual respondent name. No email, contact, upload, signature or portal action was performed.',
 'Technical evidence, source checkpoint and fixture ZIP are demonstration transfers, not automatically permissible bid attachments. Owner must select allowed PDF/DOCX/XLSX/image response files and review public-information/confidentiality implications before any authorized submission.'])]
md='# Phase 6 proposal working draft — NOT A BID\n\n'+'\n\n'.join('## '+title+'\n\n'+'\n\n'.join(paras) for title,paras in sections)+'\n';(draft/'PHASE-6-PROPOSAL-DRAFT.md').write_text(md)
doc=Document();doc.styles['Normal'].font.size=Pt(10);doc.sections[0].footer.paragraphs[0].text='DRAFT — [OWNER LEGAL ENTITY REQUIRED] — RFP #20261404054'
for i,(title,paras) in enumerate(sections):
 if i:doc.add_page_break()
 doc.add_heading(title,0 if i==0 else 1)
 for p in paras:doc.add_paragraph(p)
docpath=out/'RFP20261404054_OWNER-REQUIRED_PROPOSAL-DRAFT.docx';doc.save(docpath)
fontdir=Path(reportlab.__file__).parent/'fonts'
pdfmetrics.registerFont(TTFont('Vera',str(fontdir/'Vera.ttf')));pdfmetrics.registerFont(TTFont('VeraBold',str(fontdir/'VeraBd.ttf')))
styles=getSampleStyleSheet()
for name in ['BodyText','Normal']:styles[name].fontName='Vera'
for name in ['Heading1','Title']:styles[name].fontName='VeraBold'
story=[];pages=[]
for i,(title,paras) in enumerate(sections):
 if i:story.append(PageBreak())
 story.append(Paragraph(escape(title),styles['Title' if i==0 else 'Heading1']));story.append(Spacer(1,12))
 for p in paras:story.extend([Paragraph(escape(p),styles['BodyText']),Spacer(1,12)])
def footer(c,d):
 pages.append(c.getPageNumber());c.saveState();c.setFont('Vera',8);c.drawString(36,24,'DRAFT / OWNER LEGAL ENTITY REQUIRED — RFP #20261404054');c.drawRightString(575,24,str(c.getPageNumber()));c.restoreState()
pdfpath=out/'RFP20261404054_OWNER-REQUIRED_PROPOSAL-DRAFT.pdf';SimpleDocTemplate(str(pdfpath),topMargin=42,bottomMargin=45,leftMargin=42,rightMargin=42).build(story,onFirstPage=footer,onLaterPages=footer);assert max(pages)<=25
manifest={'draftOnly':True,'readyToSubmit':False,'submitted':False,'signed':False,'actualVendorPricesEntered':False,'officialExhibitSha256':price['sha256'],'officialBytesUnchanged':official.read_bytes()==original.read_bytes(),'officialTemplateZerosAreNotVendorPrices':True,'internalCostInputsBlank':True,'formulasStructurallyVerified':True,'spreadsheetApplicationRecalculationExecuted':False,'pdfPages':max(pages),'executiveSummaryPDFPages':1,'docxPageLayoutValidated':False,'libraries':{p:importlib.metadata.version(p) for p in ['openpyxl','python-docx','reportlab']},'files':[{'file':str(p.relative_to(root)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [official,path,docpath,pdfpath]]}
(out/'proposal-pricing-verification.json').write_text(json.dumps(manifest,indent=2)+'\n');print(json.dumps(manifest,indent=2))
